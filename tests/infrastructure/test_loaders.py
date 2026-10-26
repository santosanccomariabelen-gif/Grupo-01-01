import json
from datetime import date

from docx import Document

from agent.infrastructure.ingestion.loaders import (
    csv_to_markdown,
    html_to_markdown,
    load_corpus,
    pdf_text_to_markdown,
    to_markdown,
)


def test_csv_becomes_a_markdown_table():
    markdown = csv_to_markdown("plazos", "tipo,dias\nfraude,10\n")
    assert markdown == "# plazos\n\n| tipo | dias |\n| --- | --- |\n| fraude | 10 |\n"


def test_html_keeps_headings_and_lists_but_drops_navigation():
    html = "<nav>Inicio</nav><h1>Canales</h1><h2>App</h2><p>Disponible 24 horas.</p><ul><li>Bloqueo</li></ul><footer>pie</footer>"
    markdown = html_to_markdown(html)
    assert markdown.splitlines()[:6] == ["# Canales", "", "## App", "", "Disponible 24 horas.", ""]
    assert "- Bloqueo" in markdown
    assert "Inicio" not in markdown and "pie" not in markdown


def test_html_tables_become_markdown_tables():
    html = "<h2>Comisiones</h2><table><tr><th>Concepto</th><th>Monto</th></tr><tr><td>Retiro</td><td>S/ 7.50</td></tr></table>"
    assert html_to_markdown(html).splitlines()[:5] == [
        "## Comisiones", "", "| Concepto | Monto |", "| --- | --- |", "| Retiro | S/ 7.50 |",
    ]


def test_docx_tables_keep_their_place(tmp_path):
    path = tmp_path / "fees.docx"
    document = Document()
    document.add_heading("Servicios", level=1)
    table = document.add_table(rows=2, cols=2)
    for row_index, row in enumerate([["Servicio", "Comisión"], ["Carta", "S/ 25"]]):
        for column_index, text in enumerate(row):
            table.cell(row_index, column_index).text = text
    document.add_paragraph("Las constancias digitales no tienen costo.")
    document.save(path)

    assert to_markdown(path).splitlines()[:6] == [
        "## Servicios", "", "| Servicio | Comisión |", "| --- | --- |", "| Carta | S/ 25 |", "",
    ]
    assert "Las constancias digitales no tienen costo." in to_markdown(path)


def test_docx_headings_follow_the_same_levels(tmp_path):
    path = tmp_path / "policy.docx"
    document = Document()
    document.add_heading("Política", level=0)
    document.add_heading("Canales", level=1)
    document.add_paragraph("App y teléfono.")
    document.save(path)

    assert to_markdown(path).splitlines()[:5] == ["# Política", "", "## Canales", "", "App y teléfono."]


def test_pdf_sections_are_detected_by_numbering():
    text = "Reglamento\n1. Objeto\nEste reglamento aplica\na todos los canales.\n3.1 Plazo para presentar\nHasta 60 días.\n4111 1111 no es un título"
    markdown = pdf_text_to_markdown(text)
    assert markdown.splitlines() == [
        "# Reglamento", "",
        "## 1. Objeto", "",
        "Este reglamento aplica a todos los canales.", "",
        "### 3.1 Plazo para presentar", "",
        "Hasta 60 días. 4111 1111 no es un título",
    ]


def test_load_corpus_reads_the_manifest(tmp_path):
    (tmp_path / "faq.md").write_text("# FAQ\n\n## Pregunta\n\nRespuesta.\n", encoding="utf-8")
    manifest = [{"source": "faq.md", "doc_type": "faq", "product": None, "valid_from": "2025-01-01", "valid_until": None}]
    (tmp_path / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")

    [document] = load_corpus(tmp_path)

    assert document.source == "faq.md" and document.valid_from == date(2025, 1, 1)
    assert document.markdown.startswith("# FAQ")
